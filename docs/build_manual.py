"""Generate docs/KernelLeakGuard_User_Manual.docx.

    .venv/bin/python docs/build_manual.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import report_lib as R  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "docs" / "figures"

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
R.setup_styles(doc)
w = R.Writer(doc)

# Title page
for text, size, bold in (("KernelLeakGuard", 28, True),
                         ("User Manual", 20, True),
                         ("Hybrid static analysis and runtime verification of memory leaks "
                          "on Linux kernel early-exit paths", 13, False),
                         ("", 12, False),
                         ("BITE303P Operating Systems Lab, Fall 2026-27, VIT Vellore", 12, False),
                         ("Om Gunjan Gupta (24BIT0316) · Abhishek Kumar Singh (24BIT0309) · "
                          "Aarav Jain (24BIT0023)", 12, False),
                         ("Guide: Dr. Arivuselvan K", 12, False)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    R._font(p.add_run(text), size, bold=bold)

w.chapter_heading(1, "Introduction")
w.para("KernelLeakGuard finds memory leaks on the error paths of Linux kernel modules and "
       "checks each finding by reproducing the failure inside a virtual machine. This manual "
       "explains how to install it, run it, read its results and analyse your own module.")
w.para("What you get for every finding:")
w.bullets([
    "the function, the allocation (file and line) and the leaking return with its error code;",
    "the call whose failure leads to the leak (for example a second kmalloc() or "
    "copy_from_user());",
    "a verdict from the virtual machine: **confirmed** (kmemleak saw the leak), **not "
    "reproduced** (the failure was injected and nothing leaked) or **inconclusive** (the "
    "failure could not be injected);",
    "the evidence: fault number, stack frame, kmemleak backtrace.",
])
w.figure(FIG / "fig_architecture.png", "How the pieces fit together", 16)

w.chapter_heading(2, "System Requirements")
w.table(["Item", "Requirement"], [
    ["Operating system", "Windows 10/11 with WSL2 (Ubuntu 24.04 or newer), or native Linux x86-64"],
    ["Virtualisation", "KVM available (/dev/kvm); on Windows 11 WSL2 provides it"],
    ["CPU / RAM", "4+ cores recommended, 8 GB RAM (the guest VM uses 2 GB)"],
    ["Disk", "about 25 GB (kernel source and build)"],
    ["Network", "only for the first download of packages and kernel source"],
], "Requirements", widths=[4, 12])

w.chapter_heading(3, "Installation")
w.h2("3.1 Install WSL2 and Ubuntu (Windows only)")
w.para("Open PowerShell as administrator and run `wsl --install -d Ubuntu`, restart, then "
       "open the Ubuntu app from the Start menu and choose a user name and password.")
w.h2("3.2 Install packages")
w.para("In the Ubuntu terminal:")
w.code("""sudo apt update
sudo apt install -y build-essential clang llvm llvm-dev lld libclang-dev \\
  qemu-system-x86 flex bison bc libelf-dev libssl-dev dwarves cpio busybox-static \\
  python3-venv python3-pip git curl xz-utils rsync kmod
sudo usermod -aG kvm $USER""")
w.para("Close the terminal and run `wsl --shutdown` in Windows PowerShell so the group change "
       "takes effect. Check with `ls -l /dev/kvm` and `id` (the kvm group must be listed).")
w.h2("3.3 Get the project and build the components")
w.code("""cd ~
git clone <repository-url> KernelLeakGuard      # or copy the project folder here
cd KernelLeakGuard
./vm/build_kernel.sh        # downloads and builds Linux 6.18.55 (10-20 min the first time)
./vm/build_initramfs.sh     # guest root filesystem with klg-agent
make -C analyzer            # static analyser -> analyzer/build/klg-analyze
./scripts/setup_venv.sh     # Python packages for dashboard and documents""")
w.para("Keep the project inside the Linux file system (for example /home/<user>), not under "
       "/mnt/c: the kernel build needs a case-sensitive file system and runs much faster there. "
       "From Windows the folder is reachable at \\\\wsl$\\Ubuntu\\home\\<user>\\KernelLeakGuard.")

w.chapter_heading(4, "Running KernelLeakGuard")
w.h2("4.1 The complete pipeline")
w.code("python3 -m klg all")
w.para("This builds the 32 benchmark modules and their LLVM IR, runs the analyser in six "
       "configurations, boots the VM twice (fault-point sweep, then confirmation) and prints "
       "the evaluation. It takes about 3-4 minutes. Expected output (numbers may vary slightly):")
w.code("""[klg] built 32 modules and their LLVM IR in 23.2s
[klg] static (full): 19 warnings in 32 modules
...
[klg] runtime verification: {'confirmed': 15, 'inconclusive': 3, 'not-reproduced': 1}
| Approach | TP | FP | FN | Precision | Recall | F1 |
| static/full | 16 | 1 | 1 | 0.94 | 0.94 | 0.94 |
| hybrid      | 15 | 0 | 2 | 1.00 | 0.88 | 0.94 |
| klg/combined | 16 | 0 | 1 | 1.00 | 0.94 | 0.97 |""")
w.h2("4.2 Individual steps")
w.table(["Command", "What it does"], [
    ["python3 -m klg build", "build corpus modules, LLVM bitcode, klg-agent and initramfs"],
    ["python3 -m klg analyze [--all-variants]", "static analysis (full, or also the ablation variants)"],
    ["python3 -m klg verify [--skip-sweep]", "runtime verification in the VM (reuse the last sweep with --skip-sweep)"],
    ["python3 -m klg evaluate", "precision/recall against the ground truth, triage table"],
    ["python3 -m klg show", "one line per finding with its verdict and reason"],
    ["python3 -m klg fix", "generate fix patches for confirmed leaks and validate them "
     "(compile, re-analyse, VM re-run); patches in results/fixes/patches"],
    ["python3 -m klg real [--only m1,m2]", "analyse the real Linux modules of "
     "vm/real_modules.fragment and verify their initialisation in the VM (results/real)"],
    ["python3 -m klg seed [--only name]", "reintroduce a known leak into a real kernel module "
     "(seeds/) and run detect -> confirm -> fix on it (results/seeded)"],
    [".venv/bin/python -m klg.figures", "regenerate the charts in docs/figures"],
    ["scripts/repeat_runs.sh N", "run the pipeline N times to check reproducibility"],
], "Command reference", widths=[6.5, 9.5])
w.h2("4.3 The dashboard")
w.code(".venv/bin/streamlit run dashboard/app.py")
w.para("Open http://localhost:8501 in a Windows browser. Tabs:")
w.bullets([
    ("Overview: ", "counts of warnings and verdicts, static vs hybrid precision, pipeline summary."),
    ("Findings: ", "filterable table of all findings. Selecting one shows the source with the "
     "allocation (blue), the failing call (orange) and the leaking exit (red), the suggested "
     "fix, and the runtime tests with kmemleak backtraces."),
    ("Evaluation: ", "comparison table and precision/recall chart; triage of warnings."),
    ("Run pipeline: ", "runs any step and streams its output into the page."),
])

w.chapter_heading(5, "Reading the Results")
w.table(["Verdict", "Meaning", "What to do"], [
    ["confirmed", "The predicted failure was injected and kmemleak reported an object "
     "allocated at the reported line as unreferenced.", "Fix the bug; the evidence shows how to reproduce it."],
    ["not reproduced", "The failure was injected but nothing leaked; the static warning is "
     "most likely a false positive (e.g. a correct rollback loop).", "Review briefly; usually no action."],
    ["inconclusive", "No injectable failure leads to the path, or the failure cannot happen "
     "in the test state.", "Review the warning manually."],
], "Verdicts", widths=[2.8, 8.2, 5])
w.table(["Field", "Meaning"], [
    ["alloc", "file and line of the allocation that leaks"],
    ["exit / return_value", "the return statement on the leaking path and the value it returns"],
    ["triggers", "calls whose failure leads onto the path (the last one is the primary)"],
    ["confidence", "high: the function frees the object elsewhere but not on this error path; "
     "medium: ownership is transferred elsewhere, or never released; low: other cases"],
    ["category", "missing-release-on-error-path, ownership-not-transferred-on-error-path, "
     "never-released, missing-release-on-other-path"],
], "Fields of a static finding (results/static/full/<module>.json)", widths=[4, 12])
w.para("Output files: results/static/<variant>/*.json (findings), results/runtime/"
       "verification.json (verdicts and evidence), results/runtime/sweep|confirm/out/ (raw "
       "agent logs, kmemleak reports, guest console), results/evaluation.json (metrics).")

w.chapter_heading(6, "Analysing Your Own Module")
w.para("1. Put the module source into a directory with a Kbuild file (`obj-m += mymod.o`) and "
       "build it against the guest kernel:")
w.code("make -C ~/klg-build/linux-6.18.55 O=$HOME/klg-build/kbuild LLVM=1 M=$PWD modules")
w.para("2. Generate bitcode and analyse it:")
w.code("""python3 -c "from klg import build; from pathlib import Path; print(build.emit_ir(Path('mymod.o').resolve(), Path('.')))"
analyzer/build/klg-analyze mymod.bc -o mymod.json
python3 -m json.tool mymod.json | less""")
w.para("3. For runtime verification, copy the module into corpus/ (or another directory "
       "passed to the verifier). Code that runs in module_init is verified automatically "
       "(init mode). For other code, expose the routine through a debugfs file with the "
       "harness macros in corpus/klg_harness.h (`KLG_TRIGGER(fn)`), then run "
       "`python3 -m klg all`.")
w.para("4. Unknown allocators or release functions can be added without recompiling via "
       "`klg-analyze --spec extra.json`, for example "
       "`{\"allocators\": [\"my_alloc\"], \"deallocators\": {\"my_free\": 0}}`.")

w.chapter_heading(7, "Troubleshooting")
w.table(["Problem", "Solution"], [
    ["/dev/kvm not accessible", "Add yourself to the kvm group (sudo usermod -aG kvm $USER) and "
     "restart WSL; without KVM the VM falls back to slow emulation."],
    ["Module build: unresolved symbols", "Run ./vm/build_kernel.sh first (it creates Module.symvers)."],
    ["VM run failed", "Read results/runtime/<run>/console.log and out/job.log; a kernel panic "
     "message there shows what went wrong."],
    ["All verdicts inconclusive", "Check that scripts/faddr2line works: "
     "~/klg-build/linux-6.18.55/scripts/faddr2line corpus/<m>.ko <func>+0x10"],
    ["Commands fail from Windows PowerShell", "Run them inside the Ubuntu terminal, or with "
     "wsl -d Ubuntu --cd ~/KernelLeakGuard -- <command>."],
    ["Dashboard not reachable", "Use http://localhost:8501; if the port is busy add "
     "--server.port 8502."],
], "Troubleshooting", widths=[5, 11])

R.page_numbers(doc)
out = ROOT / "docs" / "KernelLeakGuard_User_Manual.docx"
doc.save(out)
print(f"wrote {out}")
