"""Shared python-docx helpers for KernelLeakGuard documents (VIT report rules:
Times New Roman, headings 16/14/12 bold, body 12, 1.5 spacing, justified,
figure captions below figures, table captions above tables)."""
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Consolas"


def _font(run, size=12, bold=None, italic=None, name=FONT, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    for lvl, size in ((1, 16), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), FONT)
        for t in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(t)) is not None:
                del rfonts.attrib[qn(t)]
        st.paragraph_format.space_before = Pt(12 if lvl > 1 else 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True


class Writer:
    """Appends formatted content to a document body."""

    def __init__(self, doc: Document):
        self.doc = doc
        self.fig = {}
        self.tab = {}
        self.chapter = 0

    # -- structure -----------------------------------------------------------
    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def chapter_heading(self, number: int | None, title: str):
        self.chapter = number or 0
        text = f"CHAPTER {number}: {title.upper()}" if number else title.upper()
        h = self.doc.add_heading(text, level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # "Page break before" instead of a break paragraph: a break paragraph
        # spills onto an empty page when the previous page is exactly full.
        h.paragraph_format.page_break_before = True
        return h

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    # -- text ----------------------------------------------------------------
    def para(self, text="", bold_prefix: str | None = None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             size=12, italic=False):
        p = self.doc.add_paragraph()
        p.alignment = align
        if bold_prefix:
            _font(p.add_run(bold_prefix), size, bold=True)
        self._rich(p, text, size, italic)
        return p

    def _rich(self, p, text, size=12, italic=False):
        """Minimal inline markup: `code` and **bold**."""
        import re
        for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
            if not part:
                continue
            if part.startswith("`"):
                _font(p.add_run(part[1:-1]), size - 1, name=MONO)
            elif part.startswith("**"):
                _font(p.add_run(part[2:-2]), size, bold=True, italic=italic)
            else:
                _font(p.add_run(part), size, italic=italic)

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            _font(p.add_run(f"{i}. " if numbered else "•  "), 12)
            if isinstance(it, tuple):
                _font(p.add_run(it[0]), 12, bold=True)
                self._rich(p, it[1])
            else:
                self._rich(p, it)

    def code(self, text, size=9):
        for line in text.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.line_spacing = 1.0
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.left_indent = Cm(0.5)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F3F3F1")
            p._p.get_or_add_pPr().append(shd)
            _font(p.add_run(line if line else " "), size, name=MONO)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # -- figures and tables --------------------------------------------------
    def figure(self, path: Path, caption: str, width_cm=16.0):
        n = self.fig.get(self.chapter, 0) + 1
        self.fig[self.chapter] = n
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(c.add_run(f"Figure {self.chapter}.{n}: {caption}"), 11, bold=True)
        return f"Figure {self.chapter}.{n}"

    def next_figure(self):
        return f"Figure {self.chapter}.{self.fig.get(self.chapter, 0) + 1}"

    def next_table(self):
        return f"Table {self.chapter}.{self.tab.get(self.chapter, 0) + 1}"

    def table(self, header, rows, caption, widths=None, size=10):
        n = self.tab.get(self.chapter, 0) + 1
        self.tab[self.chapter] = n
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = True
        _font(c.add_run(f"Table {self.chapter}.{n}: {caption}"), 11, bold=True)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            _font(cell.paragraphs[0].add_run(str(h)), size, bold=True)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "E7E6E1")
            cell._tc.get_or_add_tcPr().append(shd)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.line_spacing = 1.0
                self._rich(para, str(v), size)
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing = 1.0
                    para.paragraph_format.space_after = Pt(2)
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return f"Table {self.chapter}.{n}"


# -- fields, footer, TOC -------------------------------------------------------

def _field(paragraph, instr: str, placeholder=""):
    r1 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc)
    r4 = paragraph.add_run(placeholder)
    _font(r4, 12)
    r5 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r5._r.append(fc)


def page_numbers(doc: Document):
    for s in doc.sections:
        s.different_first_page_header_footer = True
        p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(p, "PAGE", "1")
        for r in p.runs:
            _font(r, 11)


def toc_field(paragraph):
    _field(paragraph, 'TOC \\o "1-2" \\h \\z \\u',
           "Right-click here and choose “Update Field” to build the table of contents.")


def update_fields_on_open(doc: Document):
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def remove_from(doc: Document, start_el):
    """Delete start_el and every following body element (keeps sectPr)."""
    body = doc.element.body
    el = start_el
    while el is not None:
        nxt = el.getnext()
        if el.tag != qn("w:sectPr"):
            body.remove(el)
        el = nxt


GUIDE = "Dr. Arivuselvan K"
GUIDE_COVER = "Dr. ARIVUSELVAN K"     # the template's cover writes the name in capitals


def set_guide(doc: Document):
    """Replace the template's guide name on the cover and in the declaration."""
    for p in doc.paragraphs:
        t = p.text
        if "VIVEKANANDA" in t or "Vivekananda" in t:
            new = t.replace("Dr. G.N. VIVEKANANDA", GUIDE_COVER).replace("Dr. G.N. Vivekananda", GUIDE)
            set_runs(p, new)


def tighten_cover(doc: Document):
    """Drop blank paragraphs on the template's cover page so that a long
    project title does not push the last cover line onto the next page."""
    set_guide(doc)
    paras = doc.paragraphs
    try:
        start = next(i for i, p in enumerate(paras) if p.text.strip() == "Submitted by")
        end = next(i for i, p in enumerate(paras) if i > start and p.text.strip().startswith("Fall"))
    except StopIteration:
        return
    for p in paras[start:end]:
        if not p.text.strip() and not p._p.xpath(".//pic:pic") and not p._p.xpath('.//w:br[@w:type="page"]'):
            p._p.getparent().remove(p._p)
    # Blank filler lines after the signature would spill onto an empty page.
    paras = doc.paragraphs
    sig = next((i for i, p in enumerate(paras) if p.text.strip().startswith("Signature of Student")), None)
    toc = next((i for i, p in enumerate(paras) if p.text.strip() == "TABLE OF CONTENTS"), None)
    if sig is not None and toc is not None:
        for p in paras[sig + 1:toc]:
            if not p.text.strip():
                p._p.getparent().remove(p._p)
    # The template relies on the cover filling page 1; make the break explicit.
    for p in doc.paragraphs:
        if p.text.strip() in ("DECLARATION", "TABLE OF CONTENTS"):
            p.paragraph_format.page_break_before = True


def set_runs(paragraph, text):
    """Replace a paragraph's text keeping the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
