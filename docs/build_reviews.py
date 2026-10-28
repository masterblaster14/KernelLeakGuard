"""Generate the Project Proposal and the Project Design & Progress Report
(Review-I) on the VIT cover page, in plain language.

    .venv/bin/python docs/build_reviews.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
import report_lib as R  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DOCS, FIG, RES = ROOT / "docs", ROOT / "docs" / "figures", ROOT / "results"
TITLE = R.TITLE
TEAM = [("Om Gunjan Gupta", "24BIT0316"), ("Abhishek Kumar Singh", "24BIT0309"),
        ("Aarav Jain", "24BIT0023")]
GUIDE = "Dr. Arivuselvan K, Associate Professor, SCORE"


def cover(kind: str) -> tuple[Document, R.Writer]:
    doc = Document(DOCS / "templates" / "OS_Project_Report_Template.docx")
    R.setup_styles(doc)
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "PROJECT REPORT":
            R.set_runs(p, kind)
        elif t == "TITLE OF THE PROJECT":
            R.set_runs(p, TITLE)
            for r in p.runs:
                r.font.size = R.Pt(17)
                r.font.color.rgb = R.RGBColor(0, 0, 0)
    rows = [p for p in doc.paragraphs if p.text.strip().startswith("Student Name")]
    for p, (n, reg) in zip(rows, TEAM):
        R.set_runs(p, f"{n}\t\t\t\t{reg}")
        for r in p.runs:
            r.font.color.rgb = R.RGBColor(0, 0, 0)
    R.tighten_cover(doc)
    decl = next(p for p in doc.paragraphs if p.text.strip() == "DECLARATION")
    R.remove_from(doc, decl._p)
    return doc, R.Writer(doc)


def finish(doc, name):
    R.page_numbers(doc)
    out = DOCS / name
    doc.save(out)
    print(f"wrote {out}")


class Flow:
    """Numbered sections (1., 1.1) without a page break before each one."""

    def __init__(self, writer):
        self.w = writer

    def section(self, number, title):
        self.w.chapter = number
        return self.w.h2(f"{number}. {title}")

    def sub(self, title):
        return self.w.h3(title)

    def shot(self, name: str, caption: str, max_w=16.0, max_h=19.5):
        """A figure scaled to fit the page (tall screenshots get narrower)."""
        path = FIG / name
        if not path.exists():
            print(f"  (missing {name}, skipped)")
            return
        w, h = Image.open(path).size
        self.w.figure(path, caption, min(max_w, max_h * w / h))

    def __getattr__(self, name):
        return getattr(self.w, name)


EXAMPLE = """\
static int sensor_setup(void)
{
        char *rx, *tx, *name;
        int ret;

        rx = kmalloc(RX_SIZE, GFP_KERNEL);       /* take memory #1 */
        if (!rx)
                return -ENOMEM;
        tx = kmalloc(TX_SIZE, GFP_KERNEL);       /* take memory #2 */
        if (!tx) {
                ret = -ENOMEM;
                goto err_rx;
        }
        name = kasprintf(GFP_KERNEL, "sensor%d", 0);
        if (!name) {                             /* something failed ...      */
                ret = -ENOMEM;
                goto err_rx;                     /* BUG: should be err_tx     */
        }
        ...
err_tx:
        kfree(tx);                               /* give back memory #2 */
err_rx:
        kfree(rx);                               /* give back memory #1 */
        return ret;
}"""

REFS = [
    "Wang, W. “MLEE: Effective Detection of Memory Leaks on Early-Exit Paths in OS "
    "Kernels.” USENIX Annual Technical Conference (ATC), 2021, pp. 31–45.",
    "Linux kernel documentation: “Kernel Memory Leak Detector (kmemleak).” "
    "https://docs.kernel.org/dev-tools/kmemleak.html",
    "Linux kernel documentation: “Fault injection capabilities infrastructure.” "
    "https://docs.kernel.org/fault-injection/fault-injection.html",
    "Google. “syzkaller: an unsupervised coverage-guided kernel fuzzer.” "
    "https://github.com/google/syzkaller",
    "Emamdoost, N., Wu, Q., Lu, K., McCamant, S. “Detecting Kernel Memory Leaks in "
    "Specialized Modules with Ownership Reasoning (K-MELD).” NDSS, 2021.",
    "Lattner, C., Adve, V. “LLVM: A Compilation Framework for Lifelong Program Analysis "
    "and Transformation.” CGO, 2004.",
    "Bellard, F. “QEMU, a Fast and Portable Dynamic Translator.” USENIX ATC, 2005.",
    "Cppcheck: a tool for static C/C++ code analysis. https://cppcheck.sourceforge.io/",
    "Silberschatz, A., Galvin, P. B., Gagne, G. Operating System Concepts, 10th ed. Wiley, 2018.",
    "Corbet, J., Rubini, A., Kroah-Hartman, G. Linux Device Drivers, 3rd ed. O’Reilly, 2005.",
]


def references(w):
    w.section(w.w.chapter + 1 if isinstance(w.w.chapter, int) else 99, "References")
    for i, r in enumerate(REFS, 1):
        p = w.para(f"[{i}] {r}", align=R.WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = R.Cm(0.9)
        p.paragraph_format.first_line_indent = R.Cm(-0.9)


# ============================================================== PROPOSAL ====
doc, w0 = cover("PROJECT PROPOSAL")
w0.chapter_heading(None, "Project Proposal")
w = Flow(w0)

w.section(1, "Introduction")
w.para(
    "Every program needs memory while it works, and it is supposed to give that memory back "
    "when it is done. The Linux kernel — the core of the operating system that runs Android "
    "phones, most web servers, cloud machines and many embedded devices — is no different. "
    "Kernel code asks for memory with functions such as `kmalloc()` and must return it with "
    "`kfree()`. If a piece of kernel code forgets to return memory, that memory is gone until "
    "the machine is restarted. This is called a **memory leak**.")
w.para(
    "Leaks in normal applications are annoying; leaks in the kernel are worse. The kernel has "
    "no garbage collector and nobody cleans up after it when a program exits, so every leaked "
    "byte stays lost. A machine that is meant to run for months slowly loses memory, gets "
    "slower and may finally crash. If an attacker can trigger the leaky code again and again, "
    "the leak even becomes a way to bring the machine down (a denial-of-service attack).")
w.para(
    "This project builds **KernelLeakGuard**, a tool that finds such leaks automatically, "
    "proves that each one is real by reproducing it inside a running Linux kernel, and "
    "suggests a fix. This proposal explains the problem, what we want to achieve, how we "
    "plan to do it, which tools we will use and how the work is shared among the three of us.")

w.section(2, "Problem Statement")
w.para(
    "Kernel functions often stop early when something goes wrong: memory runs out, a device "
    "does not answer, a user passes a bad value. These early stops are called **error exits** "
    "or **early-exit paths**. On every such exit the function must give back everything it "
    "took so far. Programmers usually do this with a ladder of clean-up labels at the end of "
    "the function. The example below shows how easy it is to get this wrong: when the third "
    "step fails, the code jumps to the wrong label and the memory taken in the second step is "
    "never returned.")
w.code(EXAMPLE, size=9)
w.para(
    "Such bugs are hard to find for three reasons. First, error exits run only when something "
    "fails, which almost never happens during normal testing, so the bug can sit in the code "
    "for years. Second, the kernel is enormous (tens of millions of lines), so reading it by "
    "hand is impossible. Third, existing automatic tools either read the code and produce "
    "many false alarms that a person must check one by one, or run the code and only find "
    "leaks on the paths they happen to reach.")
w.para(
    "**Problem statement:** design and build a tool that automatically finds memory leaks on "
    "the error exits of Linux kernel code, confirms each suspected leak with evidence from a "
    "real running kernel instead of leaving it to a human, and proposes a checked fix.",
    align=R.WD_ALIGN_PARAGRAPH.JUSTIFY)

w.section(3, "Motivation")
w.para(
    "Memory leaks on error exits are one of the most common bug types fixed in the Linux "
    "kernel; the research tool MLEE alone reported 120 new ones. Each is small, but together "
    "they matter, because the kernel runs on billions of devices that are rarely restarted. "
    "We also found the topic a very good fit for an Operating Systems lab: it touches memory "
    "management, kernel modules, error handling, virtual machines and the debugging tools "
    "built into Linux itself, and it lets us work with real kernel code instead of a "
    "simulation.")
w.bullets([
    ("Reliability: ", "servers and phones run for weeks; a slow leak eventually causes "
     "out-of-memory failures."),
    ("Security: ", "a leak that a user can trigger repeatedly can exhaust the kernel's memory "
     "(denial of service)."),
    ("Developer time: ", "a tool that only says “maybe a leak here” still leaves "
     "hours of manual checking; proof attached to each report saves that time."),
    ("Learning: ", "the project makes us work directly with the kernel's memory allocator, "
     "module loading and fault injection — core OS topics."),
])

w.section(4, "Objectives")
w.bullets([
    ("Find: ", "build a static analyser that reads compiled kernel module code (LLVM IR) and "
     "points out memory that is not given back on some error exit, together with the "
     "failing call that leads there."),
    ("Prove: ", "build a runtime checker that boots our own Linux kernel in a virtual machine, "
     "makes exactly that call fail using the kernel's fault-injection feature, and asks the "
     "kernel's own leak detector (kmemleak) whether memory was really lost."),
    ("Decide: ", "label every warning as confirmed, not reproduced, or inconclusive, and keep "
     "the evidence (kernel messages, stack traces) with it."),
    ("Fix: ", "generate a patch for each confirmed leak and check it: it must compile, the "
     "warning must disappear and the leak must be gone when tested again."),
    ("Measure: ", "build a test set of kernel modules with known leaks and known safe code, "
     "and compare our tool with simpler approaches and an existing free tool."),
    ("Show: ", "provide a command-line tool, a web dashboard and clear documentation so that "
     "others can use and check the results."),
], numbered=True)

w.section(5, "Scope")
w.para("To keep the project achievable in one semester, we set clear limits.")
w.table(["In scope", "Out of scope"], [
    ["Memory leaks on error exits of Linux kernel modules (C code)",
     "Other bug types (use-after-free, races, deadlocks), except when we happen to find them"],
    ["Kernel memory functions: kmalloc, kzalloc, vmalloc, kmem_cache, kstrdup and similar",
     "Leaks of other resources (file handles, locks, reference counts) in general"],
    ["Our own test modules and real modules that need no special hardware",
     "Drivers that need real hardware to load"],
    ["Linux 6.18 running in QEMU on a laptop (WSL2 on Windows)",
     "Running on physical machines or other operating systems"],
    ["Error exits reached while a module loads (and test triggers in our modules)",
     "Code reached only by mounting disks or special system calls (future work)"],
], "Project scope", widths=[8, 8])

w.section(6, "Background and Literature Survey")
w.sub("6.1 Background")
w.para(
    "**Kernel memory allocation.** Kernel code gets memory from the slab allocator "
    "(`kmalloc()`, `kzalloc()`, `kmem_cache_alloc()`) or from `vmalloc()` for large areas, "
    "and returns it with `kfree()`, `vfree()` or `kmem_cache_free()`. Unlike user programs, "
    "there is no automatic clean-up.")
w.para(
    "**kmemleak.** Linux has a built-in leak detector. It remembers every allocation and, when "
    "asked, scans the kernel's memory for pointers to each one. An allocation that nobody "
    "points to any more is reported as an “unreferenced object”, together with the "
    "call stack that allocated it. This is our source of proof.")
w.para(
    "**Fault injection.** Linux can deliberately make allocations fail, for testing. The "
    "file `/proc/self/fail-nth` lets a program say “make the n-th allocation of my next "
    "system call fail”. This lets us steer the kernel onto exactly the error exit we want "
    "to test.")
w.para(
    "**LLVM.** The kernel can be compiled with clang into LLVM IR, a simple, well-defined form "
    "of the program that is much easier to analyse than C source. Our analyser works on it.")
w.sub("6.2 Existing work")
w.table(["Approach", "Example", "Strength", "Weakness"], [
    ["Read the code (static analysis)", "MLEE (2021)",
     "Finds leaks on paths that never run; found 120 real leaks", "Many false alarms; a person "
     "must check each warning"],
    ["Ownership reasoning", "K-MELD (2021)", "Understands who owns memory",
     "Static only; no proof"],
    ["Pattern checkers", "Coccinelle, Smatch, cppcheck", "Fast, easy to run",
     "Only find simple patterns; miss leaks that cross functions or branches"],
    ["Run the code (fuzzing)", "syzkaller + kmemleak", "Real evidence",
     "Only finds what it happens to reach; no explanation of the cause"],
    ["Fail every allocation", "fail-nth sweeps", "Systematic", "Very many runs; slow"],
], "Existing approaches", widths=[3.6, 3.2, 4.4, 4.8], size=9.5)
w.sub("6.3 Gap we address")
w.para(
    "No existing tool combines the two sides: static tools find candidates but cannot prove "
    "them, and dynamic tools prove leaks but cannot aim. Our idea is to let the static "
    "analyser **aim** and the running kernel **prove**: the analyser says exactly which "
    "allocation leaks, on which exit, and which failing call leads there; the virtual machine "
    "then makes only that call fail and kmemleak gives the verdict.")

w.section(7, "Proposed Method")
w.para(
    "KernelLeakGuard works in three stages, shown in the architecture figure below.")
w.shot("fig_architecture.png", "Planned architecture of KernelLeakGuard", 15.5)
w.sub("7.1 Stage 1 — Find (static analysis)")
w.bullets([
    "Compile each kernel module with clang into LLVM IR, using the kernel's own build flags.",
    "For every memory allocation, follow the pointer through copies, function calls and "
    "local structures.",
    "Search every path from the allocation to each return statement. If a path neither "
    "frees the memory nor hands it over to someone else (for example by storing it in a "
    "global list), report a possible leak.",
    "Remember which call failed on that path (the “trigger”) and the error that is "
    "returned.",
])
w.sub("7.2 Stage 2 — Prove (runtime verification)")
w.bullets([
    "Build our own Linux kernel with kmemleak and fault injection switched on, and boot it in "
    "QEMU (a virtual machine), so failures cannot harm the laptop.",
    "Inside the virtual machine, a small agent program loads the module and fails the "
    "allocations one by one (1st, 2nd, 3rd, ...) to learn which failure number belongs to "
    "which line of code.",
    "Then it fails exactly the trigger found in stage 1, three times, and asks kmemleak for "
    "unreferenced objects.",
    "If an object allocated at the reported line is left behind, the warning is "
    "**confirmed**; if the path ran but nothing leaked, it is **not reproduced**; if the "
    "failure could not be triggered, it is **inconclusive**.",
])
w.sub("7.3 Stage 3 — Fix and report")
w.bullets([
    "For a confirmed leak, add the missing free at the right place (usually by re-using an "
    "existing clean-up line), compile the module, analyse it again and re-run the test.",
    "Show everything in a command-line report and a web dashboard with the source code, the "
    "failing call and the kernel's evidence.",
])
w.shot("fig_flowchart.png", "Planned flow of the runtime verification", 10.5)

w.section(8, "Operating System Concepts Involved")
w.table(["OS concept", "How the project uses it"], [
    ["Memory management", "Kernel allocators (slab, kmalloc, vmalloc, kmem_cache), freeing "
     "memory, and how kmemleak scans memory to find lost objects"],
    ["Kernel modules", "Loading and unloading modules (insmod, finit_module), module "
     "initialisation and exit code"],
    ["Error handling", "Error codes such as -ENOMEM, clean-up ladders with goto, error "
     "exits"],
    ["Virtualization", "Running a separate Linux kernel in QEMU, sharing files with the "
     "host (9p), serial console"],
    ["System calls and /proc, debugfs", "fail-nth in /proc, kmemleak and fault-injection "
     "controls in debugfs"],
    ["Linux system programming", "A C agent that loads modules and drives fault injection; "
     "shell scripts for the guest"],
    ["Performance analysis", "Counting how many kernel runs each method needs; precision and "
     "recall"],
], "Operating-system concepts", widths=[4.2, 11.8])

w.section(9, "Tools and Software")
w.table(["Tool", "Purpose"], [
    ["Linux 6.18 (LTS) source", "The kernel we build, test and analyse"],
    ["Clang / LLVM 21", "Compile the kernel and produce LLVM IR; library for our analyser"],
    ["C++17", "The static analyser (klg-analyze)"],
    ["C", "Test modules and the in-guest agent (klg-agent)"],
    ["QEMU 10", "Virtual machine that runs our kernel"],
    ["Python 3", "Orchestration, verdicts, evaluation, reports"],
    ["Streamlit, pandas, matplotlib", "Web dashboard and charts"],
    ["BusyBox", "Tiny root file system inside the virtual machine"],
    ["cppcheck", "Existing tool we compare against"],
    ["Git", "Version history and backup"],
    ["Ubuntu on WSL2 (Windows 11)", "Development machine"],
], "Tools and software", widths=[5, 11])
w.para("**Hardware:** one laptop (8 or more CPU cores, 16 GB RAM, about 10 GB free disk). "
       "No special devices are needed.")

w.section(10, "Evaluation Plan")
w.para(
    "To know whether the tool really works, we will write our own set of small kernel "
    "modules where we know the right answer: some contain leaks (marked as such in the "
    "source) and some contain tricky but correct code that a naive tool would flag. We will "
    "then count:")
w.bullets([
    ("True positives: ", "real leaks the tool reports."),
    ("False positives: ", "safe code the tool wrongly reports."),
    ("False negatives: ", "real leaks the tool misses."),
    ("Precision and recall: ", "the share of reports that are correct, and the share of leaks "
     "that are found."),
    ("Cost: ", "how many times the kernel code had to be run."),
])
w.para(
    "We will compare: our static analysis alone, the complete tool, a naive "
    "“missing-kfree” check, failing every allocation blindly, random fault "
    "injection, and cppcheck. If time allows we will also run the tool on real Linux modules.")

w.section(11, "Expected Outcomes")
w.bullets([
    "A working tool that finds error-exit leaks and proves or refutes each warning.",
    "A benchmark of kernel modules with known answers and a comparison with other methods.",
    "Generated and checked fix patches.",
    "A dashboard, a user manual, a final report and a presentation.",
])

w.section(12, "Work Plan")
w.table(["Week", "Work", "Result"], [
    ["1", "Study kmemleak, fault injection, MLEE; set up WSL2, clang, QEMU", "Environment ready"],
    ["2", "Build Linux 6.18 with debugging features; boot it in QEMU", "Kernel boots in the VM"],
    ["3", "Write the benchmark modules and the guest agent", "Modules load in the VM"],
    ["4", "Static analyser: allocations, paths, frees", "First warnings"],
    ["5", "Ownership/escape handling, function summaries", "Fewer false alarms"],
    ["6", "Fault-point mapping and targeted confirmation", "First confirmed leak"],
    ["7", "Verdicts, evidence, symbolising stack traces", "Review-I prototype"],
    ["8", "Baselines (exhaustive, random, cppcheck), metrics", "Comparison tables"],
    ["9", "Fix patch generation and validation; dashboard", "Patches and dashboard"],
    ["10", "Real kernel modules; testing; documentation", "Final report, manual, slides"],
], "Work plan", widths=[1.5, 9.5, 5])
w.para("Milestones: **Review-I** — the core pipeline runs end to end on the benchmark "
       "(about 60–70% complete). **Review-II** — complete tool, evaluation, dashboard and "
       "documentation.")

w.section(13, "Team Members and Responsibilities")
w.table(["Member", "Register no.", "Responsibilities"], [
    ["Om Gunjan Gupta", "24BIT0316", "Static analyser in C++/LLVM: allocation tracking, "
     "ownership, path search, triggers, fix locations"],
    ["Abhishek Kumar Singh", "24BIT0309", "Kernel build, virtual machine, guest init, "
     "fault-injection agent, module loading and symbol logging"],
    ["Aarav Jain", "24BIT0023", "Benchmark modules and ground truth, Python verifier and "
     "verdicts, evaluation and baselines, dashboard"],
], "Team responsibilities", widths=[4, 2.8, 9.2])
w.para("All three of us will test, write documentation and prepare the presentation "
       "together, and each will be able to explain every part during the reviews. "
       f"Guide: **{GUIDE}**.")

w.section(14, "Risks and How We Will Handle Them")
w.table(["Risk", "What we will do"], [
    ["Kernel crashes while we inject failures", "Run everything in a virtual machine; detect "
     "crashes and restart automatically"],
    ["Too many false alarms from the analyser", "Model ownership hand-over and error codes; "
     "let the runtime stage refute the rest"],
    ["kmemleak misses a leak", "Repeat each test and clear stale memory (zeroing, stack "
     "erasing) so leaks are not hidden"],
    ["The virtual machine is slow", "Use KVM acceleration when available; only fail the calls "
     "the analyser points at"],
    ["Real modules need hardware", "Choose modules that run without hardware (virtual network "
     "devices, file systems, block simulators)"],
], "Risks", widths=[6, 10])

w.chapter = 14
references(w)
finish(doc, "KernelLeakGuard_Project_Proposal.docx")


# ============================================================== REVIEW-I ====
doc, w0 = cover("PROJECT DESIGN & PROGRESS REPORT (REVIEW-I)")
w0.chapter_heading(None, "Project Design & Progress Report (Review-I)")
w = Flow(w0)

w.section(1, "Project at a Glance")
w.para(
    "KernelLeakGuard finds memory leaks on the error exits of Linux kernel code and proves "
    "them inside a real running kernel. In one sentence: **a static analyser points at the "
    "suspicious exit and the failing call that leads there; a virtual machine makes exactly "
    "that call fail; the kernel's leak detector, kmemleak, says whether memory was really "
    "lost.** This report describes the design, how the work is split, the algorithms, how far "
    "the implementation has come by Review-I, screenshots of the working prototype, and the "
    "problems we ran into.")
w.table(["Item", "Status at Review-I"], [
    ["Problem, objectives, literature", "Done (see the proposal)"],
    ["Architecture and module design", "Done"],
    ["Static analyser", "Working on the benchmark; refinements continue"],
    ["Kernel, virtual machine, fault injection", "Working"],
    ["Runtime verification and verdicts", "Working end to end"],
    ["Benchmark modules with known answers", "Done (32 modules)"],
    ["Baselines, fix patches, dashboard, real modules", "In progress for Review-II"],
], "Summary of progress", widths=[7.5, 8.5])

w.section(2, "System Architecture")
w.para(
    "The system has two halves that talk through files. On the **host** (Ubuntu in WSL2) the "
    "static analyser reads the compiled modules and writes a list of suspected leaks. In the "
    "**guest** (our own Linux 6.18 kernel inside QEMU) the agent loads each module and "
    "injects failures, and kmemleak reports what was lost. The Python orchestrator on the "
    "host connects the two: it prepares the guest's job, boots the virtual machine, reads the "
    "guest's logs back, and turns them into verdicts.")
w.shot("fig_architecture.png", "System architecture", 16)
w.para("The main pieces, in the order the data flows:")
w.bullets([
    ("Build: ", "each module is compiled twice — normally (for the VM) and into LLVM IR "
     "(for the analyser), with exactly the kernel's own compiler flags."),
    ("klg-analyze (C++): ", "finds leak candidates and writes them as JSON."),
    ("klg (Python): ", "plans which failure to inject for each candidate and writes a shell "
     "script for the guest."),
    ("QEMU guest: ", "boots our kernel, runs the script, loads modules, injects failures, runs "
     "kmemleak, writes logs to a shared folder."),
    ("Verifier: ", "reads the logs, matches leaked objects to source lines and decides each "
     "verdict."),
    ("Reports: ", "command-line summary now; dashboard and documents for Review-II."),
])

w.section(3, "Module Design")
w.para("We split the work into four modules with clear inputs and outputs, so each of us "
       "could work and test independently.")
w.table(["Module", "Owner", "Input", "Output"], [
    ["Static analyser (analyzer/)", "Om Gunjan Gupta", "LLVM IR of a module",
     "candidates.json: allocation, leaking exit, failing call, error code"],
    ["Kernel and VM (vm/, runtime/)", "Abhishek Kumar Singh", "Kernel source, modules, guest "
     "job script", "Booted guest; agent logs; kmemleak reports"],
    ["Verifier and orchestrator (klg/)", "Aarav Jain", "candidates.json, guest logs",
     "Verdicts with evidence (verification.json)"],
    ["Benchmark and evaluation (corpus/, klg/evaluate.py)", "Aarav Jain",
     "Modules with GT:LEAK / GT:SAFE marks", "Precision, recall, comparison tables"],
], "Modules", widths=[4.2, 3.3, 3.8, 4.7], size=9.5)
w.sub("3.1 Static analyser")
w.para(
    "Written in C++ on top of the LLVM 21 libraries. It first simplifies the IR (turns stack "
    "variables into plain values), then for each allocation call it collects every value "
    "that is a copy of the pointer (aliases), notes where the pointer is freed, returned or "
    "stored somewhere outside the function (ownership hand-over), and searches the paths to "
    "every return. Small summaries of helper functions (does this function free its "
    "argument? return new memory?) let it follow memory across function calls.")
w.sub("3.2 Kernel, virtual machine and agent")
w.para(
    "We build Linux 6.18.55 with clang and switch on kmemleak, the fault-injection framework "
    "(failslab, fail_function, fail_usercopy) and a few settings that keep kmemleak "
    "accurate. QEMU boots it with a tiny BusyBox file system and shares a host folder with "
    "it. The agent (`klg-agent`, written in C) loads a module or calls a test trigger while "
    "`/proc/self/fail-nth` makes the chosen allocation fail, and prints markers that the host "
    "can parse. A helper module, `klg_modwatch`, logs where modules are loaded so that stack "
    "traces can be turned back into source lines even after a module was unloaded.")
w.sub("3.3 Verifier")
w.para(
    "The Python verifier first runs a **mapping** pass (fail allocation 1, 2, 3, ... and "
    "record where each failure happened), then picks for each candidate the failure number "
    "that hits its trigger, runs those **confirmation** tests three times, and compares "
    "kmemleak's report with the candidate's allocation site.")
w.sub("3.4 Benchmark")
w.para(
    "32 small kernel modules written by us: 17 leaking allocation sites and 42 safe ones. "
    "Each allocation line carries a comment `/* GT:LEAK */` or `/* GT:SAFE */`, which is the "
    "ground truth for scoring. The safe modules contain the tricky patterns that fool naive "
    "tools: memory handed to a list, freed through a wrapper, freed in a loop, or reference "
    "counted.")
w.sub("3.5 Data formats")
w.para("We do not need a database; all stages exchange JSON files, which are easy to "
       "inspect. A shortened candidate written by the analyser looks like this:")
w.code("""{
  "function": "sensor_setup",
  "allocator": "kmalloc",
  "alloc":  {"file": "klg_l02_goto_ladder.c", "line": 32},
  "exit":   {"file": "klg_l02_goto_ladder.c", "line": 54},
  "return_value": "-ENOMEM",
  "triggers": [{"callee": "kasprintf", "line": 38, "kind": "alloc"}],
  "confidence": "high",
  "category": "missing-release-on-error-path"
}""", size=9)

w.section(4, "Algorithms and Flowcharts")
w.sub("4.1 Finding leaks (static analysis)")
w.code("""for each allocation call A in function F:
    aliases  = every value that is a copy of A's result
    events   = places where an alias is freed, returned, or stored outside F
    depth-first search over the paths from A to every return of F:
        skip the edge where A itself failed (A == NULL or IS_ERR(A))
        skip branches that contradict an earlier decision on the same condition
        stop the path at a free / hand-over event       (no leak on this path)
        at a return: report "A leaks here",
                     with the failing call that chose this path (trigger)""", size=9)
w.para(
    "In plain words: start where memory is taken, walk every possible way out of the "
    "function, and complain about any way out that neither frees the memory nor gives it "
    "to someone else. The “skip” rules remove paths that cannot really happen.")
w.sub("4.2 Proving leaks (runtime verification)")
w.code("""mapping:  for n = 1, 2, 3, ...:
              load the module with fail-nth = n (only failures inside the module count)
              record which source line the n-th failure hit; stop when nothing fails
confirm:  for each candidate:
              pick n whose failure hits the candidate's trigger
              run 3 times with fail-nth = n; then kmemleak scan
verdict:  object from the candidate's line leaked  -> confirmed
          path ran, nothing leaked                  -> not reproduced
          trigger could not be failed               -> inconclusive""", size=9)
w.shot("fig_flowchart.png", "Flowchart of the runtime verification", 10.5)

w.section(5, "Implementation Progress")
w.para(
    "By Review-I the core of the tool runs from start to finish on the benchmark: it compiles "
    "the modules, finds candidates, boots the kernel, injects the predicted failures and "
    "produces verdicts with kmemleak evidence. We estimate the project at about two thirds "
    "complete.")
w.table(["Part", "Status", "Notes"], [
    ["Kernel build with kmemleak + fault injection", "Done", "Linux 6.18.55, clang"],
    ["QEMU guest, init script, shared folder", "Done", "BusyBox, 9p share"],
    ["klg-agent (sweep / repeat modes)", "Done", "fail-nth with module filter"],
    ["klg_modwatch symbol logging", "Done", "for unloaded modules"],
    ["Benchmark: 32 modules with ground truth", "Done", "17 leaks, 42 safe sites"],
    ["Static analyser: aliases, paths, triggers", "Done", "C++/LLVM"],
    ["Ownership hand-over, summaries, branch correlation", "Done", "reduces false alarms"],
    ["Mapping and confirmation runs, verdicts", "Done", "end to end"],
    ["Baselines (exhaustive, random, cppcheck)", "In progress", "Review-II"],
    ["Fix patch generation and validation", "In progress", "Review-II"],
    ["Dashboard", "Prototype", "findings page works"],
    ["Real kernel modules", "Planned", "Review-II"],
    ["Report, manual, presentation", "In progress", "Review-II"],
], "Implementation status", widths=[7.5, 2.5, 6], size=9.5)

w.section(6, "Screenshots of the Working Prototype")
w.para("The screenshots below were taken from the running system.")
w.shot("ss_term_analyzer.png", "The static analyser's report for one benchmark module", 15)
w.shot("shot_findings.png", "Findings and verdicts on the command line (python3 -m klg show)", 16)
w.shot("ss_vm_boot.png", "Our Linux 6.18.55 kernel booting inside the QEMU virtual machine", 15)
w.shot("shot_agent.png", "The agent injecting a failure inside the guest", 16)
w.shot("ss_vm_step3_fault.png", "A failure injected while a real module (erofs) loads", 15)
w.shot("shot_kmemleak.png", "kmemleak reporting the leaked object, with its call stack", 16)
w.shot("ss_dash_findings.png", "Dashboard prototype: the list of findings with verdicts", 16)

w.section(7, "Challenges Faced")
w.para("Not everything worked the first time. These were the hardest problems and how we "
       "solved them.")
w.table(["Challenge", "What happened", "How we solved it"], [
    ["Fault injection ignored", "Failing a function by name did nothing for our calls.",
     "The kernel's probe mechanism had to be built in its ftrace-based mode; with the other "
     "mode the injection never matched our task."],
    ["Leaks hidden from kmemleak", "A real leak was found in some runs and missed in others.",
     "Old copies of pointers on reused stacks and in reused memory made kmemleak think the "
     "memory was still used. We clear stacks and zero memory on allocation and free."],
    ["Module gone before the scan", "When a module fails to load, Linux removes it before "
     "kmemleak reports, so stack traces showed only raw addresses.",
     "We wrote klg_modwatch, a tiny module that logs every module's addresses and symbols."],
    ["Random injection crashed the guest", "PID 1 also got failures and the kernel panicked.",
     "Restrict injection to the agent's own task before turning it on."],
    ["Too many false alarms", "Helpers that return error codes instead of NULL looked like "
     "leaks.", "We model IS_ERR() like a NULL check and let runtime checks refute the rest."],
    ["Stack traces to source lines", "The kernel's faddr2line script broke with clang builds.",
     "We call it in the mode that works with LLVM builds."],
    ["Slow virtual machine", "Windows security features blocked hardware virtualization "
     "(KVM) inside WSL2.", "QEMU falls back to software emulation; the tool only fails the "
     "calls the analyser points at, which keeps runs short."],
], "Challenges and solutions", widths=[3.3, 6, 6.7], size=9.5)

w.section(8, "Plan for Review-II")
w.bullets([
    "Run the baselines (fail every allocation, random failures, cppcheck) and compare.",
    "Generate fix patches and check them by compiling, re-analysing and re-testing.",
    "Finish the dashboard and write the user manual and final report.",
    "Try the tool on real Linux modules that need no hardware.",
    "Repeat the experiments several times to check that results are stable.",
])

w.chapter = 8
references(w)
finish(doc, "KernelLeakGuard_Review1_Design_Progress.docx")
