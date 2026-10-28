"""Generate docs/KernelLeakGuard_Presentation.pptx (16:9, 20 slides).

    .venv/bin/python docs/build_ppt.py
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "docs" / "figures"
RES = ROOT / "results"
ev = json.loads((RES / "evaluation.json").read_text())
ver = json.loads((RES / "runtime" / "verification.json").read_text())
static = {p.stem: json.loads(p.read_text()) for p in sorted((RES / "static" / "full").glob("*.json"))}
A, GT = ev["approaches"], ev["ground_truth"]
n_warn = sum(len(r["candidates"]) for r in static.values())
cnt = {s: sum(1 for v in ver["verdicts"] if v["status"] == s)
       for s in ("confirmed", "not-reproduced", "inconclusive")}
_fx = RES / "fixes" / "fixes.json"
_fixes = [f for f in (json.loads(_fx.read_text()) if _fx.exists() else [])
          if f["status"] in ("validated", "runtime-validated", "failed", "skipped")]
n_fix, n_fix_ok = len(_fixes), sum(1 for f in _fixes if f["status"] == "validated")

INK, INK2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x89, 0x87, 0x81)
BLUE, ORANGE = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34)
PANEL = RGBColor(0xF3, 0xF3, 0xF1)
FONT = "Segoe UI"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def text(slide, x, y, w, h, s, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font=FONT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, font
    return tb


def bullets(slide, x, y, w, h, items, size=18, gap=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, body = it if isinstance(it, tuple) else (None, it)
        r = p.add_run()
        r.text = "•  "
        r.font.size, r.font.color.rgb, r.font.name = Pt(size), BLUE, FONT
        if head:
            r = p.add_run()
            r.text = head
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, INK, FONT
        r = p.add_run()
        r.text = body
        r.font.size, r.font.color.rgb, r.font.name = Pt(size), INK2 if head else INK, FONT
    return tb


def slide(title, num=True):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.9), title, 30, True)
    if num:
        text(s, Inches(11.8), Inches(7.0), Inches(1.3), Inches(0.4), str(len(prs.slides)), 11,
             color=MUTED, align=PP_ALIGN.RIGHT)
        text(s, Inches(0.6), Inches(7.0), Inches(8), Inches(0.4),
             "KernelLeakGuard · BITE303P Operating Systems Lab", 11, color=MUTED)
    return s


def code(slide, x, y, w, h, s, size=13):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = PANEL
    box.line.fill.background()
    box.adjustments[0] = 0.04
    tb = text(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), h - Inches(0.2),
              s.split("\n"), size, font="Consolas")
    return tb


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def table(slide, x, y, w, rows, col_w, size=14, row_h=0.42):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
    t = shape.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = str(v)
            para = c.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(size)
            para.runs[0].font.name = FONT
            para.runs[0].font.bold = i == 0
            para.runs[0].font.color.rgb = INK
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xE7, 0xE6, 0xE1) if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
    return t


def stat(slide, x, y, value, label, color=BLUE):
    text(slide, x, y, Inches(3.8), Inches(0.9), value, 40, True, color)
    text(slide, x, y + Inches(0.85), Inches(3.8), Inches(0.8), label, 14, color=INK2)


# 1 Title -----------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.25))
band.fill.solid(); band.fill.fore_color.rgb = BLUE; band.line.fill.background()
text(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(1.2), "KernelLeakGuard", 54, True)
text(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.4),
     "Detecting and Verifying Linux Kernel Memory Leaks", 30, color=INK2)
text(s, Inches(0.9), Inches(4.3), Inches(11.5), Inches(1.6), [
    "Om Gunjan Gupta (24BIT0316)   ·   Abhishek Kumar Singh (24BIT0309)   ·   Aarav Jain (24BIT0023)",
    "Guide: Dr. Arivuselvan K, Associate Professor, SCORE",
    "BITE303P Operating Systems Lab · Fall 2026-27 · VIT Vellore"], 17, color=INK2)

# 2 Problem ---------------------------------------------------------------------
s = slide("The problem: leaks on early-exit paths")
bullets(s, Inches(0.6), Inches(1.4), Inches(6.2), Inches(5), [
    "Kernel code checks every operation and bails out early with an error code.",
    "Each early exit must undo everything done so far; a missing kfree() leaks memory "
    "until reboot (no garbage collector, no process exit).",
    "These paths run only when an allocation or device operation fails, so ordinary "
    "testing almost never executes them.",
    "Static tools report candidates, but many are safe: ownership moves to lists, "
    "callers, callbacks.",
], 19)
code(s, Inches(7.1), Inches(1.5), Inches(5.7), Inches(3.9),
     "int initialize(void)\n{\n    char *a = kmalloc(100, GFP_KERNEL);\n    if (!a)\n"
     "        return -ENOMEM;\n    char *b = kmalloc(200, GFP_KERNEL);\n    if (!b)\n"
     "        return -ENOMEM;   // a leaks\n    ...\n    kfree(b);\n    kfree(a);\n    return 0;\n}", 15)
text(s, Inches(7.1), Inches(5.5), Inches(5.7), Inches(0.5),
     "Leaks only when the second kmalloc() fails", 14, color=MUTED)

# 3 Motivation ------------------------------------------------------------------
s = slide("Why it matters")
stat(s, Inches(0.8), Inches(1.6), "120", "previously unknown early-exit leaks found\nin Linux by MLEE (USENIX ATC 2021)")
stat(s, Inches(4.8), Inches(1.6), "weekly", "kmemleak reports filed by syzbot,\nGoogle's kernel fuzzing service", ORANGE)
stat(s, Inches(8.8), Inches(1.6), "0", "times an error path runs in a normal\ntest run that never fails an allocation", BLUE)
bullets(s, Inches(0.8), Inches(4.0), Inches(11.8), Inches(3), [
    ("Impact: ", "memory pressure, degraded performance, instability; an attacker who can "
     "trigger the failure repeatedly can exhaust kernel memory (denial of service)."),
    ("Gap: ", "static warnings are not proof; dynamic tools only see the paths they happen "
     "to execute and give a backtrace without the root cause."),
], 19)

# 4 Objectives ------------------------------------------------------------------
s = slide("Objectives")
bullets(s, Inches(0.6), Inches(1.4), Inches(12), Inches(5.5), [
    "Detect missing releases on early-exit paths with an LLVM-based static analyser "
    "(alias, ownership-transfer, interprocedural and path-correlation analysis).",
    "Predict, for each warning, the failing call that leads onto the leaking path.",
    "Build an isolated QEMU/KVM test kernel with kmemleak and per-task fault injection.",
    "Automatically inject the predicted failure and decide confirmed / not reproduced / "
    "inconclusive from kernel evidence.",
    "Evaluate on a benchmark with ground truth: static vs hybrid vs dynamic-only, ablation.",
    "Deliver a CLI and a dashboard with source context, evidence and suggested fixes.",
], 20, 12)

# 5 Literature ------------------------------------------------------------------
s = slide("Literature survey")
table(s, Inches(0.6), Inches(1.4), Inches(12.1), [
    ["Work", "Technique", "Strength", "Limitation"],
    ["MLEE (ATC'21)", "LLVM IR, early-exit vs normal path", "120 new kernel leaks", "static only, manual triage"],
    ["Hector (DSN'13)", "compare error paths", "no specification needed", "false positives"],
    ["K-MELD (NDSS'21)", "ownership reasoning", "custom allocators", "static only"],
    ["HERO (Sec'21)", "precise function pairing", "many confirmed bugs", "no runtime evidence"],
    ["FIFUZZ (Sec'20)", "context-sensitive fault injection", "deep error paths", "crashes, not leaks"],
    ["syzkaller + kmemleak", "fuzzing + fail-nth + kmemleak", "real evidence", "only reached paths, no root cause"],
], [2.5, 3.6, 2.9, 3.1], 14)

# 6 Gap -------------------------------------------------------------------------
s = slide("Research gap and our idea")
bullets(s, Inches(0.6), Inches(1.4), Inches(12), Inches(2.8), [
    "Static detectors cover all paths but every warning needs manual triage.",
    "Dynamic tools give evidence only for executed paths, without saying which exit forgot what.",
    "No surveyed system uses the static analysis to drive targeted fault injection per "
    "warning and feeds kernel evidence back into the verdict.",
], 20)
code(s, Inches(0.6), Inches(4.4), Inches(12.1), Inches(1.8),
     "static warning  ──►  predicted failing call  ──►  fail exactly that call in a VM  ──►  kmemleak\n"
     "                                                                                        │\n"
     "      confirmed  /  not reproduced  /  inconclusive   + reproduction + backtrace  ◄─────┘", 15)

# 7 Architecture ----------------------------------------------------------------
s = slide("System architecture")
picture(s, FIG / "fig_architecture.png", Inches(0.9), Inches(1.2), w=Inches(11.5))

# 8 Static analyser -------------------------------------------------------------
s = slide("Static analyser (klg-analyze, C++ / LLVM 21)")
bullets(s, Inches(0.6), Inches(1.4), Inches(7.0), Inches(5.5), [
    ("IR: ", "re-run each module's kbuild command with -emit-llvm; mem2reg; fold "
     "configuration stubs."),
    ("Aliases: ", "casts, GEPs, PHIs, and fields of local structs/arrays."),
    ("Handled = ", "kfree & friends, or ownership transfer (store to non-local memory, "
     "return, capturing call)."),
    ("Summaries: ", "frees / captures / returns each parameter; allocation wrappers."),
    ("Path search: ", "DFS to every return, skip the alloc-failed edge and contradictory "
     "branches (correlation), loops ×2."),
    ("Report: ", "exit, returned errno, failing trigger calls, confidence."),
], 17)
code(s, Inches(7.9), Inches(1.5), Inches(4.9), Inches(4.6),
     '{\n "function": "sensor_setup",\n "allocator": "kmalloc",\n "alloc": {"line": 32},\n'
     ' "exit": {"line": 54},\n "return_value": "-ENOMEM",\n "triggers": [\n'
     '   {"callee": "kasprintf",\n    "line": 38}],\n "confidence": "high"\n}', 14)

# 9 Runtime verifier ------------------------------------------------------------
s = slide("Runtime verifier (QEMU/KVM guest, Linux 6.18.55)")
bullets(s, Inches(0.6), Inches(1.4), Inches(12), Inches(5.5), [
    ("Fault-point sweep: ", "klg-agent writes N to /proc/thread-self/fail-nth (N = 1, 2, 3 ...) "
     "and runs the routine; the stack trace of each injected fault gives N → source line."),
    ("Targeted injection: ", "for each warning, inject the N that fails its trigger call, "
     "3 times (failslab, fail_function via ftrace kprobes, fail_usercopy)."),
    ("Kmemleak: ", "wait 6 s, scan twice; attribute each unreferenced object by its "
     "allocation backtrace (faddr2line) or, for unloaded init code, by allocation size."),
    ("Trustworthy evidence: ", "no stale pointer copies: page-allocator stacks, KSTACK_ERASE, "
     "no task-stack scanning, init_on_alloc / init_on_free for slab memory."),
    ("Speed: ", f"a guest boots in ~2 s; both VM runs take about "
     f"{round((ver['timing'].get('sweep_vm_s', 0) + ver['timing'].get('confirm_vm_s', 0)) / 60 + 0.5)} min."),
], 18, 10)

# 10 OS concepts ----------------------------------------------------------------
s = slide("Operating-system concepts in the project")
table(s, Inches(0.6), Inches(1.4), Inches(12.1), [
    ["Concept", "Where"],
    ["Memory management", "slab allocator, kmalloc caches, vmalloc, ownership, kmemleak mark & scan"],
    ["Error handling / faults", "errno propagation, goto cleanup, failslab / fail_function / fail_usercopy"],
    ["Processes & system calls", "per-task fault state (fail-nth), task stacks, write / finit_module / pread"],
    ["Device management", "loadable modules, misc devices, platform drivers, devm memory"],
    ["Synchronisation", "mutex-protected lists, RCU deferred free, kref reference counting"],
    ["Virtualisation", "QEMU + KVM, virtio-9p sharing, initramfs boot"],
    ["Performance analysis", "analysis time, VM time, executions, precision / recall"],
], [3.4, 8.7], 15, 0.5)

# 11 Corpus ---------------------------------------------------------------------
s = slide("Benchmark corpus (dataset)")
stat(s, Inches(0.8), Inches(1.5), str(GT["modules"]), "kernel modules\n(17 leaky, 15 leak-free)")
stat(s, Inches(4.3), Inches(1.5), str(GT["leak_sites"]), "leaking allocation sites", ORANGE)
stat(s, Inches(7.8), Inches(1.5), str(GT["safe_sites"]), "safe allocation sites")
bullets(s, Inches(0.8), Inches(3.6), Inches(11.8), Inches(3.4), [
    "Patterns from real kernel bugs: wrong goto label, copy_from_user failure, loops, "
    "wrappers, registration failures, conditional allocation, kmem_cache, vmalloc.",
    "Tricky safe code: list ownership, caller-owned fields, devm, kref, RCU, factories, "
    "correlated conditions, rollback loops.",
    "Ground truth in the source: every allocation is annotated /* GT:LEAK */ or /* GT:SAFE */.",
], 18)

# 12 Implementation & team ------------------------------------------------------
s = slide("Implementation and module division")
table(s, Inches(0.6), Inches(1.4), Inches(12.1), [
    ["Member", "Modules"],
    ["Om Gunjan Gupta (24BIT0316)", "Static analyser (C++/LLVM): IR preparation, alias/escape, summaries, path search"],
    ["Abhishek Kumar Singh (24BIT0309)", "Guest kernel & VM, klg-agent fault-injection driver (C), kmemleak protocol"],
    ["Aarav Jain (24BIT0023)", "Corpus & ground truth, Python verifier & evaluation, Streamlit dashboard"],
], [3.8, 8.3], 15, 0.55)
bullets(s, Inches(0.6), Inches(4.2), Inches(12), Inches(2.5), [
    "C++17 + LLVM/Clang 21 · C (static) · Python 3 · Linux 6.18.55 LTS · QEMU/KVM on WSL2",
    "Single command: python3 -m klg all  →  build, analyse, verify in VM, evaluate",
], 17)

# 13 Results --------------------------------------------------------------------
s = slide("Results: detection quality")
rows = [["Approach", "TP", "FP", "FN", "Prec.", "Rec."]]
for n, label in (("cppcheck", "cppcheck 2.19"), ("static/naive", "Static, naive"),
                 ("static/full", "Static, KLG"), ("hybrid", "Hybrid (confirmed)"),
                 ("klg/combined", "KLG complete"), ("dynamic/exhaustive", "Dynamic, exhaustive"),
                 ("dynamic/random", "Dynamic, random")):
    if n not in A:
        continue
    a = A[n]
    rows.append([label, a["tp"], a["fp"], a["fn"], f"{a['precision']:.2f}", f"{a['recall']:.2f}"])
table(s, Inches(0.5), Inches(1.4), Inches(5.6), rows, [2.2, 0.6, 0.6, 0.6, 0.9, 0.7], 12, 0.44)
picture(s, FIG / "fig_precision_recall.png", Inches(6.3), Inches(1.5), w=Inches(6.8))
text(s, Inches(0.5), Inches(5.2), Inches(5.6), Inches(1.6),
     f"KernelLeakGuard: {A['klg/combined']['tp']} of {GT['leak_sites']} leaks, no false "
     f"positive; every confirmed leak has kernel evidence.", 15, color=INK2)

# 14 Verdicts + ablation ---------------------------------------------------------
s = slide("Runtime verdicts and ablation")
picture(s, FIG / "fig_triage.png", Inches(0.4), Inches(1.3), w=Inches(6.2))
picture(s, FIG / "fig_static_breakdown.png", Inches(6.8), Inches(1.3), w=Inches(6.2))
bullets(s, Inches(0.5), Inches(5.0), Inches(6.0), Inches(2), [
    (f"{cnt['confirmed']} confirmed, {cnt['not-reproduced']} not reproduced, "
     f"{cnt['inconclusive']} inconclusive ", "— no false confirmation"),
], 14, 6)
bullets(s, Inches(6.9), Inches(5.0), Inches(6.0), Inches(2), [
    ("Without ownership transfer: ", f"precision {A['static/no-escape']['precision']:.2f}; "
     f"naive: {A['static/naive']['fp']} false alarms"),
], 14, 6)

# 15 Real kernel code ---------------------------------------------------------------
s = slide("Real Linux 6.18.55 kernel code")
_rs = RES / "real" / "summary.json"
_real = json.loads(_rs.read_text()) if _rs.exists() else {}
_tri = json.loads((RES / "real" / "triage.json").read_text()) if (RES / "real" / "triage.json").exists() else []
_rev = sum(1 for t in _tri if t.get("basis") == "reviewed in source")
_crashed = sorted({c.get("where", {}).get("log", "").removeprefix("sweep_")
                   for c in _real.get("crash_details", []) if c.get("where", {}).get("log")})
_seeds = json.loads((RES / "seeded" / "summary.json").read_text()) if (RES / "seeded" / "summary.json").exists() else []
_sdet = sum(1 for s_ in _seeds if s_.get("detected"))
_sconf = [s_ for s_ in _seeds if s_.get("confirmed")]
_rc = json.loads((RES / "real" / "root_causes.json").read_text()) if (RES / "real" / "root_causes.json").exists() else {}
_sl = _rc.get("slub_sysfs_leak", {})
bullets(s, Inches(0.6), Inches(1.4), Inches(12), Inches(5.6), [
    ("Static: ", f"{_real.get('modules', 0)} modules, {_real.get('functions', 0)} functions, "
     f"{_real.get('alloc_sites', 0)} allocation sites → {_real.get('warnings', 0)} warnings "
     "(115 before general refinements: ERR_PTR, ownership of published objects, infeasible "
     "cleanup paths; benchmark unchanged)"),
    ("Runtime (VM): ", f"{len(_real.get('runtime_modules', []))} modules, "
     f"{_real.get('fault_points', 0)} fault points of module init exercised; no warning confirmed"),
    ("Re-introduced leaks: ", "; ".join(
        f"{s_['module']} {s_['function']}()"
        + (f" (real bug, fixed upstream by {s_['history']['commit']})" if s_.get("history") else " (seeded)")
        + (" → detected, confirmed by kmemleak in the VM, fixed" if s_.get("confirmed")
           else " → detected" if s_.get("detected") else " → not detected")
        for s_ in _seeds) or "—"),
    ("SLUB leak (root-caused): ", (f"init of {', '.join(_sl.get('also_in', ['erofs']))} exposed a slab-cache "
                                   f"leak in mm/slub.c when sysfs registration fails; still in mainline; our patch: "
                                   f"{_sl['stock']['unreferenced']} → {_sl['patched']['unreferenced']} "
                                   "leaked objects in the VM" if _sl else "—")),
    ("f2fs regression (runtime): ", ("kobject names leak since commit b635f2ecdb (Apr 2026, also in "
                                     "stable); our patch: "
                                     f"{sum(_rc['f2fs_sysfs_leak']['stock_leaked_names'].values())} → "
                                     f"{sum(_rc['f2fs_sysfs_leak']['patched_leaked_names'].values())} leaks"
                                     if _rc.get("f2fs_sysfs_leak") else "—")),
    ("Static finds: ", ("ntfs3 do_action() leaks an mft_inode when mi_format_new() fails; ufs "
                        "ubh_bforget() never frees the ufs_buffer_head (reported 2018, unfixed); "
                        "both in mainline, patches remove the warnings"
                        if _rc.get("ntfs3_do_action_leak") else "—")),
    ("null_blk crash (root-caused): ", (f"use-after-free of module memory: deferred fput of the "
                                        "partition-scan file runs after a failed init freed the module"
                                        if _crashed else "no crash")),
    ("Lesson: ", "static precision drops on real code; runtime verification keeps reports clean "
     "and its fault injection finds real kernel defects"),
], 15, 9)

# 16 Case study -----------------------------------------------------------------
s = slide("Case study: klg_l02 — wrong goto label")
code(s, Inches(0.5), Inches(1.4), Inches(6.3), Inches(4.9),
     "tx = kmalloc(TX_SIZE, GFP_KERNEL);   // 32\nif (!tx) {\n    ret = -ENOMEM;\n    goto err_rx;\n}\n"
     "name = kasprintf(GFP_KERNEL, ...);  // 38\nif (!name) {\n    ret = -ENOMEM;\n"
     "    goto err_rx;   // BUG: err_tx\n}\n...\nerr_tx:\n    kfree(tx);\nerr_rx:\n    kfree(rx);\n    return ret;", 13)
case = next((v for v in ver["verdicts"] if v["module"] == "klg_l02_goto_ladder"), None)
t0 = case["tests"][0] if case and case["tests"] else {}
bullets(s, Inches(7.1), Inches(1.5), Inches(5.9), Inches(5), [
    ("Static: ", "kmalloc at 32 leaks when returning -ENOMEM; trigger kasprintf (38)."),
    ("Sweep: ", f"fault #{t0.get('fail_nth')} fails {t0.get('expected_frame')}."),
    ("Confirm: ", f"injected {t0.get('fault_hits')}×, kmemleak: {len(t0.get('leaks', []))} "
     "unreferenced objects allocated at line 32."),
    ("Verdict: ", f"{case['status'] if case else ''}. Fix: goto err_tx."),
], 16, 12)

# 17 Performance ----------------------------------------------------------------
s = slide("Performance")
picture(s, FIG / "fig_cost.png", Inches(0.5), Inches(1.4), w=Inches(7.0))
bullets(s, Inches(7.8), Inches(1.5), Inches(5.2), Inches(5), [
    ("Static analysis: ", f"{A['static/full']['analysis_ms']:.0f} ms for 32 modules"),
    ("Mapping + confirmation VM runs: ", f"{ver['timing'].get('mapping_vm_s', 0):.0f} s + "
     f"{ver['timing'].get('confirm_vm_s', 0):.0f} s"),
    ("KernelLeakGuard: ", f"{A['hybrid']['executions']} executions in total (module-filtered "
     f"sweep + confirmation) vs {A['dynamic/exhaustive']['executions']} exhaustive"),
    ("Fix patches: ", f"{n_fix_ok} of {n_fix} validated (compile, re-analysis, VM re-run)"),
    ("Hardware: ", "laptop, Ryzen 7, 15 GB, WSL2 + KVM"),
], 16, 10)

# 18 Demo -----------------------------------------------------------------------
s = slide("Demonstration")
code(s, Inches(0.5), Inches(1.35), Inches(5.6), Inches(2.6),
     "cd ~/KernelLeakGuard\nbash scripts/demo_ready.sh --open\n\n"
     "python3 scripts/live_demo.py --window\n"
     "python3 scripts/live_demo.py shell --window", 14)
bullets(s, Inches(0.5), Inches(4.15), Inches(5.6), Inches(3), [
    "VM screen: our Linux 6.18.55 boots in QEMU, shown in the browser",
    "erofs loads cleanly; one allocation failed (fault 85); kmemleak reports the leaked "
    "slab cache; the kernel warns (SLUB bug)",
    "The 2010 bonding bug: kmemleak reports the leaked net_device",
    "Dashboard: findings, evaluation, real-kernel root causes, Live kernel page",
], 14, 6)
if (FIG / "ss_vm_step4_kmemleak.png").exists():
    picture(s, FIG / "ss_vm_step4_kmemleak.png", Inches(6.4), Inches(1.35), w=Inches(6.6))

# 19 Conclusion -----------------------------------------------------------------
s = slide("Conclusion and future work")
bullets(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(5.5), [
    "Static analysis finds early-exit leaks with F1 "
    f"{A['static/full']['f1']:.2f} (naive: {A['static/naive']['f1']:.2f}).",
    "Targeted fault injection + kmemleak confirms them with evidence; precision "
    f"{A['hybrid']['precision']:.2f}.",
    f"Complete tool: {A['klg/combined']['tp']}/{GT['leak_sites']} leaks, 0 false positives, "
    f"validated fix patches for {n_fix_ok}/{n_fix}.",
    "Predicting the failing call + module-filtered injection: cheaper than blind injection.",
    "kmemleak must be protected from stale pointer copies (stacks, reused slab memory).",
], 18, 12)
text(s, Inches(7.0), Inches(1.4), Inches(5.8), Inches(0.5), "Future work", 20, True)
bullets(s, Inches(7.0), Inches(2.0), Inches(5.8), Inches(5), [
    "Field-sensitive heap ownership (nested leaks)",
    "Real kernel code: 39 hardware-free modules already built into the guest",
    "Reproduce MLEE's historical leaks",
    "Loop-counter reasoning; patches for shared error labels",
    "CI integration (KernelCI)",
], 17, 8)

# 20 Thanks ---------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.25))
band.fill.solid(); band.fill.fore_color.rgb = BLUE; band.line.fill.background()
text(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.2), "Thank you", 54, True, align=PP_ALIGN.CENTER)
text(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(0.8), "Questions?", 28, color=INK2,
     align=PP_ALIGN.CENTER)
text(s, Inches(0.9), Inches(5.0), Inches(11.5), Inches(0.6),
     "Om Gunjan Gupta · Abhishek Kumar Singh · Aarav Jain", 16, color=MUTED, align=PP_ALIGN.CENTER)

out = ROOT / "docs" / "KernelLeakGuard_Presentation.pptx"
prs.save(out)
print(f"wrote {out} ({len(prs.slides)} slides)")
